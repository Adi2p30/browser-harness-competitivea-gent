"""Excel <-> DataFrame.

Program sheets: row 1 = headers; A = university, B = school, C = program name, D = "Website Link";
every column after "Website Link" is a field to fill. Other sheets (notes, empty templates) are ignored.
"""
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment

# Not stated on a college's own site (third-party rankings, derived numbers, our own columns).
NOT_ON_COLLEGE_SITE = re.compile(r"rank|cost of living|roi|breakeven|confidence|conferral", re.I)
LINK_COLUMN = re.compile(r"\blink\b|source url", re.I)


def load_frames(path: str) -> dict[str, pd.DataFrame]:
    """One DataFrame per sheet; blank cells are NaN and values keep their Excel types."""
    return pd.read_excel(path, sheet_name=None, dtype=object)


def website_column(df: pd.DataFrame) -> int | None:
    for i, col in enumerate(df.columns[:6]):
        if "website link" in str(col).casefold():
            return i
    return None


def classify_columns(df: pd.DataFrame) -> tuple[list[str], list[str], dict[str, str]]:
    """(fields to research, link columns to derive from a paired field, {skipped column: reason})."""
    fields, links, skipped = [], [], {}
    for col in df.columns[(website_column(df) or 0) + 1:]:
        if str(col).startswith("Unnamed"):
            continue
        if NOT_ON_COLLEGE_SITE.search(str(col)):
            skipped[col] = "not stated on the college's own site (ranking / derived / cost-of-living / confidence)"
        elif LINK_COLUMN.search(str(col)):
            links.append(col)
        else:
            fields.append(col)
    return fields, links, skipped


def paired_field(link_col: str, fields: list[str]) -> str | None:
    """The researched field a link column points at, e.g. '... (Link)' -> '... (Yes/No)'."""
    if "source url" in link_col.casefold():
        stem = "tuition"
    else:
        stem = re.sub(r"[\s(]*\blink\b[\s)]*", " ", link_col, flags=re.I).strip().casefold()
    return next((f for f in fields if stem and f.casefold().startswith(stem)), None)


def is_blank(value) -> bool:
    return pd.isna(value) or str(value).strip() == ""


def write_back(src: str, dest: str, frames: dict[str, pd.DataFrame],
               changed: dict[str, list[tuple[int, str, str]]]) -> None:
    """Write only the changed cells into the original workbook (formatting survives).

    Each change is (row, column, note); the note (source URL + quote) becomes a cell comment.
    """
    wb = load_workbook(src)
    for sheet, cells in changed.items():
        ws, df = wb[sheet], frames[sheet]
        for row, col, note in cells:
            cell = ws.cell(row=row + 2, column=list(df.columns).index(col) + 1, value=df.at[row, col])
            cell.comment = Comment(note, "competitive-analysis-agent", width=420, height=160)
    wb.save(dest)
