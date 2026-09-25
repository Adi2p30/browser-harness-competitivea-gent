"""Build retry batches: every Online-sheet cell the first pass left "Not found" or "UNVERIFIED".

Same job grouping as ../batches; each row keeps only its missing fields plus the first pass's note.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
ROOT = HERE.parents[2]
wb = load_workbook(ROOT / "EDIT Competitive Analysis Research 2026-2027_online_filled.xlsx")
out = HERE / "batches"
out.mkdir(exist_ok=True)
total = 0
for path in sorted((HERE.parent / "batches").glob("batch_*.json")):
    jobs = []
    for job in json.loads(path.read_text()):
        rows = []
        for r in job["rows"]:
            ws = wb[r["sheet"]]
            hdr = {c.value: c.column for c in ws[1]}
            missing = {}
            for field in r["fields"] + r["link_columns"]:
                cell = ws.cell(row=r["row"], column=hdr[field])
                if isinstance(cell.value, str) and cell.value.startswith(("Not found", "UNVERIFIED")):
                    missing[field] = (cell.comment.text if cell.comment else "")[:300]
            if missing:
                rows.append({**r, "fields": [f for f in missing if f not in r["link_columns"]],
                             "link_columns": [f for f in missing if f in r["link_columns"]], "first_pass": missing})
                total += len(missing)
        if rows:
            jobs.append({**job, "rows": rows})
    if jobs:
        (out / path.name).write_text(json.dumps(jobs, indent=1))
print(f"{total} cells to retry")
