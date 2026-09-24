import json
import logging
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "agent-workspace"))

from competitive_analysis import __main__ as cli  # noqa: E402
from competitive_analysis import agent, claude, colleges, crawler, excel_io, rcac, skills  # noqa: E402
from competitive_analysis.crawler import Page  # noqa: E402
from competitive_analysis.rcac import parse_json  # noqa: E402

CTX = agent.Context(date(2026, 9, 18), "Fall 2027", "MS in Computer Science")
PAD = " Lorem ipsum filler text." * 12  # pages under 200 chars are treated as unreadable


def finding(field, value, quote, cycle="exact", program="target"):
    return {"field": field, "value": value, "quote": quote, "cycle_match": cycle, "program_match": program}


def script(monkeypatch, pages, replies):
    """pages: {url: (text, [links])}; replies: {url: reply dict}. Returns the list of prompts sent."""
    prompts = []
    monkeypatch.setattr(agent, "fetch", lambda url: Page(pages[url][0] + PAD, [("l", u) for u in pages[url][1]], "")
                        if url in pages else None)

    def fake_llm(prompt):
        prompts.append(prompt)
        if "Pick the candidate" in prompt:
            return json.dumps(replies["resolve"])
        url = prompt.split("Page URL: ")[1].split("\n")[0]
        return json.dumps(replies[url])

    monkeypatch.setattr(agent, "call_rcac", fake_llm)
    return prompts


def test_parse_json_handles_fences_and_prose():
    assert parse_json('Sure!\n```json\n{"found": []}\n```\nDone.') == {"found": []}
    with pytest.raises(ValueError):
        parse_json("no json here")


def test_fetch_keeps_same_site_html_links_only(monkeypatch):
    class Resp:
        status_code, url = 200, "https://www.example.edu/a"
        headers = {"content-type": "text/html", "last-modified": "Mon, 01 Jan 2026"}
        content = b"x"
        text = ('<script>x</script><p>Tuition $1</p><a href="/b#top">B</a><a href="/f.pdf">pdf</a>'
                '<a href="https://admissions.example.edu/c">C</a><a href="https://other.com/d">D</a>')

    monkeypatch.setattr(crawler.requests, "get", lambda *a, **k: Resp)
    page = crawler.fetch("https://www.example.edu/a")
    assert page.text == "Tuition $1 B pdf C D"
    assert page.last_modified == "Mon, 01 Jan 2026"
    assert [u for _, u in page.links] == ["https://www.example.edu/b", "https://admissions.example.edu/c"]


@pytest.mark.parametrize("change,reason", [
    ({"quote": "The fee is $99 for everyone"}, "quote not found on page"),   # invented quote
    ({"value": "$75"}, "value not in quote"),                                 # value contradicts its quote
    ({"cycle_match": "different"}, "cycle different"),                        # last year's page
    ({"program_match": "other"}, "program other"),                            # undergraduate, not the MS
    ({"field": "Bogus"}, "unknown field"),
])
def test_rejection_reasons(change, reason):
    good = finding("Fee", "$50", "Application fee: $50.")
    assert agent._rejection(good, ["Fee"], agent._squash("Apply now. Application fee: $50. Thanks.")) is None
    assert agent._rejection({**good, **change}, ["Fee"], agent._squash("Apply now. Application fee: $50. Thanks.")) == reason


def test_stale_and_wrong_program_findings_are_dropped_and_prompt_carries_context(monkeypatch):
    prompts = script(monkeypatch,
        {"https://x.edu/": ("Deadline for Fall 2026 is Dec 1. Undergrad fee is $60.", [])},
        {"https://x.edu/": {"found": [
            finding("Deadline", "Dec 1", "Deadline for Fall 2026 is Dec 1.", cycle="different"),
            finding("Fee", "$60", "Undergrad fee is $60.", program="other")], "next_urls": []}})
    results, meta = agent.research_college("X", ["https://x.edu/"], ["Deadline", "Fee"], CTX)
    assert {f: r["status"] for f, r in results.items()} == {"Deadline": "not-found", "Fee": "not-found"}
    assert meta["rejected"] == {"cycle different": 1, "program other": 1}
    assert "2026-09-18" in prompts[0] and "Fall 2027" in prompts[0] and "MS in Computer Science" in prompts[0]


def test_field_settles_only_when_two_pages_agree_and_crawl_then_stops(monkeypatch):
    q = "Application fee: $50 for Fall 2027."
    prompts = script(monkeypatch, {
        "https://x.edu/": ("home " + q, ["https://x.edu/apply", "https://x.edu/other"]),
        "https://x.edu/apply": ("apply " + q, []),
        "https://x.edu/other": ("never fetched", [])},
        {"https://x.edu/": {"found": [finding("Fee", "$50", q)], "next_urls": ["https://x.edu/apply"]},
         "https://x.edu/apply": {"found": [finding("Fee", "$50", q)], "next_urls": []}})
    results, meta = agent.research_college("X", ["https://x.edu/"], ["Fee"], CTX)
    assert results["Fee"]["status"] == "confirmed" and results["Fee"]["value"] == "$50"
    assert meta["visited"] == ["https://x.edu/", "https://x.edu/apply"]  # kept going past first hit, stopped once settled
    assert len(prompts) == 2


def test_single_exact_source_is_accepted_but_labelled_and_unstated_needs_corroboration(monkeypatch):
    script(monkeypatch, {"https://x.edu/": ("Deadline Dec 1 2026. Fee $50.", [])},
        {"https://x.edu/": {"found": [
            finding("Deadline", "Dec 1", "Deadline Dec 1 2026.", cycle="exact"),
            finding("Fee", "$50", "Fee $50.", cycle="unstated")], "next_urls": []}})
    results, _ = agent.research_college("X", ["https://x.edu/"], ["Deadline", "Fee"], CTX)
    assert results["Deadline"]["status"] == "single-source"
    assert results["Fee"]["status"] == "unverified-cycle" and results["Fee"]["value"] is None


def test_program_specific_value_beats_general_and_exact_beats_unstated(monkeypatch):
    script(monkeypatch, {
        "https://x.edu/": ("Fee $90 general. Fee $50 for CS MS. Old $10.", ["https://x.edu/b"]),
        "https://x.edu/b": ("Fee $50 for CS MS again." , [])},
        {"https://x.edu/": {"found": [
            finding("Fee", "$90", "Fee $90 general.", program="general"),
            finding("Fee", "$50", "Fee $50 for CS MS.", program="target"),
            finding("Fee", "$10", "Old $10.", cycle="unstated")], "next_urls": ["https://x.edu/b"]},
         "https://x.edu/b": {"found": [finding("Fee", "$50", "Fee $50 for CS MS again.")], "next_urls": []}})
    results, _ = agent.research_college("X", ["https://x.edu/"], ["Fee"], CTX)
    assert (results["Fee"]["value"], results["Fee"]["status"]) == ("$50", "confirmed")


def test_conflict_goes_to_resolver_and_null_choice_leaves_blank(monkeypatch):
    for choice, status, value in [(2, "resolved-conflict", "$75"), (None, "conflict-unresolved", None)]:
        script(monkeypatch, {"https://x.edu/": ("Fee $50 domestic. Fee $75 international.", [])},
            {"https://x.edu/": {"found": [
                finding("Fee", "$50", "Fee $50 domestic."), finding("Fee", "$75", "Fee $75 international.")],
                "next_urls": []},
             "resolve": {"choice": choice, "reason": "r"}})
        results, _ = agent.research_college("X", ["https://x.edu/"], ["Fee"], CTX)
        assert (results["Fee"]["status"], results["Fee"]["value"]) == (status, value)


def test_max_pages_and_unreadable_pages(monkeypatch):
    pages = {f"https://x.edu/{i}": (f"page {i}", [f"https://x.edu/{i + 1}"]) for i in range(20)}
    pages["https://x.edu/0"] = ("tiny", [])
    monkeypatch.setattr(agent, "fetch", lambda url: Page("tiny", [], "") if url == "https://x.edu/0" else None)
    _, meta = agent.research_college("X", ["https://x.edu/0", "https://x.edu/9"], ["Fee"], CTX)
    assert meta["unreadable"] == ["https://x.edu/0", "https://x.edu/9"]
    prompts = script(monkeypatch, pages, {u: {"found": [], "next_urls": []} for u in pages})
    agent.research_college("X", ["https://x.edu/1"], ["Fee"], CTX, max_pages=3)
    assert len(prompts) == 3


HEADERS = ["University Name", "School / College Name", "Program Name", "Website Link", "Tuition Per Credit",
           "Experiential Learning (Yes/No)", "Experiential Learning Link", "Ranking (QS)", "Confidence"]


def _workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Programs"
    ws.append(HEADERS)
    ws.append(["ALPHA U - WEST LAFAYETTE", None, "MS Supply Chain", None, None, None, None, None, None])
    ws.append(["Nowhere College", None, None, None, None, None, None, None, None])  # no link -> skipped
    ws["A1"].font = Font(bold=True)
    wb.create_sheet("Notes").append(["not a program sheet"])
    wb.save(path)


@pytest.fixture
def clean_logging(monkeypatch):
    monkeypatch.setattr(cli, "pick_models", lambda requested: ["rcac"])  # no API keys in tests
    monkeypatch.setattr(crawler, "BROWSER_FIRST", False)
    root = logging.getLogger()
    before = list(root.handlers)
    yield
    for h in root.handlers[:]:
        if h not in before:
            h.close()
            root.removeHandler(h)
    crawler.PAGE_DIR = rcac.LLM_LOG = claude.LLM_LOG = None


def test_classify_columns_and_link_pairing():
    import pandas as pd
    df = pd.DataFrame(columns=HEADERS + ["Unnamed: 9"])
    fields, links, skipped = excel_io.classify_columns(df)
    assert fields == ["Tuition Per Credit", "Experiential Learning (Yes/No)", "Ranking (QS)"]  # schools cite rankings
    assert links == ["Experiential Learning Link"]
    assert set(skipped) == {"Confidence"}
    assert excel_io.paired_field("Experiential Learning Link", fields) == "Experiential Learning (Yes/No)"
    assert excel_io.paired_field("Tuition Source URL", fields) == "Tuition Per Credit"
    assert excel_io.paired_field("Curriculum (link)", fields) is None


def test_links_for_matches_sheet_college_and_program_with_purdue_alias():
    assert colleges.links_for("Online MSGSCM", "PURDUE UNIVERSITY - WEST LAFAYETTE", "anything")
    assert colleges.links_for("Online MSGSCM", "Nowhere College") == []


def test_cli_fills_verified_cells_links_and_logs_everything(monkeypatch, tmp_path, capsys, clean_logging):
    src = tmp_path / "in.xlsx"
    _workbook(src)
    monkeypatch.setitem(colleges._INDEX, ("Programs", "alpha u", ""), ["https://alpha.edu/"])
    seen = {}

    def fake_research(college, urls, fields, ctx, max_pages, max_depth, models, verify):
        seen.update(college=college, fields=fields, ctx=ctx, models=models, verify=verify)
        crawler.log.info("FETCH %s", urls[0])  # stands in for the real crawler's logging
        return ({"Tuition Per Credit": {"value": "$500", "status": "confirmed", "source": "https://alpha.edu/cost",
                                        "quote": "Tuition is $500 per credit.", "candidates": []},
                 "Experiential Learning (Yes/No)": {"value": "Yes", "status": "single-source", "source": "https://alpha.edu/exp",
                                                    "quote": "Students complete a capstone.", "candidates": []}},
                {"visited": urls, "unreadable": [], "rejected": {}})

    monkeypatch.setattr(cli, "research_college", fake_research)
    argv = ["prog", str(src), "--cycle", "2026-2027 academic year", "--as-of", "2026-09-18", "--leave-blank"]
    monkeypatch.setattr(sys, "argv", argv)
    cli.main()

    out = load_workbook(tmp_path / "in_filled.xlsx")["Programs"]
    assert seen["college"] == "ALPHA U - WEST LAFAYETTE"  # the workbook's own name is what is researched
    assert seen["fields"] == ["School / College Name", "Tuition Per Credit", "Experiential Learning (Yes/No)",
                              "Ranking (QS)"]  # no links or derived columns
    assert (seen["models"], seen["verify"]) == (("rcac",), True)
    assert "MS Supply Chain" in seen["ctx"].program and "Programs" in seen["ctx"].program
    assert (seen["ctx"].today, seen["ctx"].cycle) == (date(2026, 9, 18), "2026-2027 academic year")
    assert [c.value for c in out[2]] == ["ALPHA U - WEST LAFAYETTE", None, "MS Supply Chain", "https://alpha.edu/", "$500",
                                         "Yes", "https://alpha.edu/exp", None, None]
    assert [c.value for c in out[3]][3:] == [None] * 6
    assert "https://alpha.edu/cost" in out["E2"].comment.text and "Tuition is $500" in out["E2"].comment.text
    assert out["A1"].font.bold
    assert load_workbook(src)["Programs"]["E2"].value is None  # input untouched

    logs = tmp_path / "in_logs"
    text = (logs / "run.log").read_text()
    assert "CELL Programs!E2 | Tuition Per Credit | None -> '$500'" in text
    assert "CELL Programs!D2 | Website Link" in text and "CELL Programs!G2 | Experiential Learning Link" in text
    assert "SKIP column 'Confidence'" in text and "NO LINK for Programs!3 Nowhere College" in text
    assert "FETCH https://alpha.edu/" in text
    report = json.loads((logs / "report.json").read_text())
    assert report["context"]["cycle"] == "2026-2027 academic year"
    assert "1 row(s) skipped for lack of a link" in capsys.readouterr().out

    # Same cycle resumes from the report without researching again; a new cycle starts fresh.
    seen.clear()
    cli.main()
    assert not seen
    monkeypatch.setattr(sys, "argv", argv[:3] + ["2027-2028 academic year"] + argv[4:])
    cli.main()
    assert seen["ctx"].cycle == "2027-2028 academic year"


def test_cli_leaves_unverified_fields_blank_and_supports_sheet_and_limit(monkeypatch, tmp_path, capsys, clean_logging):
    src = tmp_path / "in.xlsx"
    _workbook(src)
    monkeypatch.setitem(colleges._INDEX, ("Programs", "alpha u", ""), ["https://alpha.edu/"])
    calls = []
    monkeypatch.setattr(cli, "research_college", lambda *a: calls.append(a) or (
        {"Tuition Per Credit": {"value": None, "status": "conflict-unresolved", "source": None, "quote": None, "candidates": []}},
        {"visited": [], "unreadable": [], "rejected": {}}))
    monkeypatch.setattr(sys, "argv", ["prog", str(src), "--cycle", "c", "--sheet", "Programs", "--limit", "1", "--leave-blank"])
    cli.main()
    assert len(calls) == 1
    assert load_workbook(tmp_path / "in_filled.xlsx")["Programs"]["E2"].value is None
    assert "LEFT BLANK Programs!2 'Tuition Per Credit': conflict-unresolved" in (tmp_path / "in_logs" / "run.log").read_text()
    assert "4 researched field(s) left blank" in capsys.readouterr().out  # school, tuition, experiential, ranking


def test_fetch_logs_every_link_and_saves_page_text(monkeypatch, tmp_path, caplog):
    class Resp:
        status_code, url = 200, "https://www.example.edu/a"
        headers = {"content-type": "text/html"}
        content = b"x"
        text = '<p>Hello</p><a href="/b">B</a><a href="https://other.com/d">D</a>'

    monkeypatch.setattr(crawler.requests, "get", lambda *a, **k: Resp)
    monkeypatch.setattr(crawler, "PAGE_DIR", tmp_path)
    with caplog.at_level(logging.DEBUG, logger="competitive_analysis.crawler"):
        crawler.fetch("https://www.example.edu/a")
    assert "link: B | https://www.example.edu/b" in caplog.text
    assert "link dropped (other site): https://other.com/d" in caplog.text
    assert "Hello" in next(tmp_path.glob("*.txt")).read_text()


def test_name_normalisation_bridges_the_two_workbooks():
    n = colleges._norm
    assert n("Indiana University - Bloomington") == n("INDIANA UNIVERSITY BLOOMINGTON")
    assert n("University of Texas - Austin") == n("UNIVERSITY OF TEXAS AT AUSTIN")
    assert n("The Ohio State University") == n("OHIO STATE UNIVERSITY")
    assert n("Purdue University") == n("PURDUE UNIVERSITY - WEST LAFAYETTE")
    assert n("Purdue University") != n("PURDUE UNIVERSITY GLOBAL")  # different school: must not match


def test_same_value_wording_is_not_a_conflict_but_different_dates_are():
    k = agent._same_value_key
    assert k("30 credit hours") == k("30") == k("30.0 credits")
    assert k("$728.00") == k("$728") == k("728 USD per credit")
    assert k("Jun 15") != k("Dec 15") and k("Jun 15") == k("jun 15")
    assert k("33") != k("3.00 per course")


def test_credit_hours_worded_differently_settle_without_resolver(monkeypatch):
    script(monkeypatch, {"https://x.edu/": ("Total 30 credit hours.", ["https://x.edu/b"]),
                         "https://x.edu/b": ("Program: 30 credits.", [])},
        {"https://x.edu/": {"found": [finding("Credits", "30 credit hours", "Total 30 credit hours.")], "next_urls": ["https://x.edu/b"]},
         "https://x.edu/b": {"found": [finding("Credits", "30", "Program: 30 credits.")], "next_urls": []}})
    results, _ = agent.research_college("X", ["https://x.edu/"], ["Credits"], CTX)
    assert results["Credits"]["status"] == "confirmed"  # no "Pick the candidate" call was scripted, so none was made


def test_cli_starts_from_the_workbooks_own_website_link(monkeypatch, tmp_path, clean_logging):
    src = tmp_path / "in.xlsx"
    _workbook(src)
    wb = load_workbook(src)
    wb["Programs"]["D3"] = "https://nowhere.edu/program"  # a link the user (or link discovery) put in the sheet
    wb.save(src)
    starts = []
    monkeypatch.setattr(cli, "research_college", lambda college, urls, *a: starts.append((college, urls)) or ({}, {}))
    monkeypatch.setattr(sys, "argv", ["prog", str(src), "--cycle", "c"])
    cli.main()
    assert starts == [("Nowhere College", ["https://nowhere.edu/program"])]


def test_apply_links_fills_blank_cells_only_and_rejects_bad_links(monkeypatch, tmp_path, clean_logging):
    from competitive_analysis import apply_links
    src = tmp_path / "in.xlsx"
    _workbook(src)
    wb = load_workbook(src)
    ws = wb["Programs"]
    ws.append(["Kept U", None, "MS", "https://kept.edu/", None, None, None, None, None])  # already linked: never overwritten
    ws.append(["Dead U", None, "MS", None, None, None, None, None, None])
    ws.append(["Guess U", None, "MS", None, None, None, None, None, None])
    wb.save(src)
    found = tmp_path / "links.json"
    found.write_text(json.dumps([
        {"sheet": "Programs", "row": 2, "url": "https://alpha.edu/ms", "confidence": "high", "program_on_page": "MS Supply Chain", "note": "n"},
        {"sheet": "Programs", "row": 3, "url": "https://nowhere.edu/ms", "confidence": "medium", "program_on_page": "MS", "note": ""},
        {"sheet": "Programs", "row": 4, "url": "https://kept.edu/new", "confidence": "high", "program_on_page": "MS", "note": ""},
        {"sheet": "Programs", "row": 5, "url": "https://dead.edu/x", "confidence": "high", "program_on_page": "MS", "note": ""},
        {"sheet": "Programs", "row": 6, "url": "https://guess.edu/x", "confidence": "low", "program_on_page": None, "note": "unsure"}]))
    monkeypatch.setattr(apply_links, "check", lambda url: {"https://dead.edu/x": "HTTP 404", "https://nowhere.edu/ms": "blocked"}.get(url, "ok"))
    monkeypatch.setattr(sys, "argv", ["prog", str(src), str(found)])
    apply_links.main()

    out = load_workbook(src)["Programs"]
    assert [out.cell(r, 4).value for r in range(2, 7)] == ["https://alpha.edu/ms", "https://nowhere.edu/ms", "https://kept.edu/", None, None]
    assert "discovered by web research (high" in out["D2"].comment.text and "loads" in out["D2"].comment.text
    assert "not verified" in out["D3"].comment.text  # blocked: kept but flagged
    assert (tmp_path / "in_before_links.xlsx").exists()
    assert load_workbook(tmp_path / "in_before_links.xlsx")["Programs"]["D2"].value is None
    text = (tmp_path / "in_logs" / "apply_links.log").read_text()
    assert "REJECTED (HTTP 404)" in text and "NOT WRITTEN (low confidence)" in text


def test_cli_fills_every_cell_no_matter_what(monkeypatch, tmp_path, clean_logging):
    src = tmp_path / "in.xlsx"
    _workbook(src)
    wb = load_workbook(src)
    wb["Programs"].title = "Online Programs"
    wb.create_sheet("Residential").append(HEADERS)
    wb["Residential"].append(["Res U", None, "MS", "https://res.edu/", None, None, None, None, None])
    wb["Online Programs"].append(["Colleges not ranked on US news as of 09/07/25", None, None, None, None, None, None, None, None])
    wb.save(src)
    monkeypatch.setitem(colleges._INDEX, ("Online Programs", "alpha u", ""), ["https://alpha.edu/", "https://alpha.edu/ms"])
    calls = []
    cand = {"value": "$510", "quote": "Tuition $510 per credit.", "cycle_match": "unstated", "program_match": "target",
            "url": "https://alpha.edu/cost", "model": "claude"}

    def fake_research(college, urls, fields, ctx, *rest):
        calls.append((college, urls))
        return ({"School / College Name": {"value": "Alpha School", "status": "single-source", "source": "https://alpha.edu/",
                                           "quote": "Alpha School", "candidates": [], "verified_by": "claude", "model": "claude"},
                 "Tuition Per Credit": {"value": None, "status": "unverified-cycle", "source": None, "quote": None,
                                        "candidates": [cand], "verify_rejected": []},
                 "Experiential Learning (Yes/No)": {"value": None, "status": "not-found", "source": None, "quote": None,
                                                    "candidates": []}},
                {"visited": ["https://alpha.edu/"], "unreadable": [], "rejected": {}})

    monkeypatch.setattr(cli, "research_college", fake_research)
    monkeypatch.setattr(sys, "argv", ["prog", str(src), "--online"])
    cli.main()
    assert calls == [("ALPHA U - WEST LAFAYETTE", ["https://alpha.edu/", "https://alpha.edu/ms"])]  # every known link, online only
    out = load_workbook(tmp_path / "in_filled.xlsx")
    row = [c.value for c in out["Online Programs"][2]]
    assert row == ["ALPHA U - WEST LAFAYETTE", "Alpha School", "MS Supply Chain", "https://alpha.edu/", "UNVERIFIED: $510",
                   cli.NOT_FOUND, cli.NOT_FOUND, cli.NOT_FOUND, cli.NOT_PUBLISHED]
    assert "verified by claude" in out["Online Programs"]["B2"].comment.text
    assert "NOT VERIFIED" in out["Online Programs"]["E2"].comment.text
    assert "https://alpha.edu/" in out["Online Programs"]["F2"].comment.text  # pages checked
    assert [c.value for c in out["Online Programs"][3]][1:] == [cli.NO_LINK, None] + [cli.NO_LINK] * 5 + [cli.NOT_PUBLISHED]
    assert [c.value for c in out["Online Programs"][4]][1:] == [None] * 8  # note row untouched
    assert out["Residential"]["E2"].value is None

    # Placeholders count as blank, so a rerun with --retry-unfilled researches them again.
    calls.clear()
    monkeypatch.setattr(sys, "argv", ["prog", str(tmp_path / "in_filled.xlsx"), "--online", "--inplace", "--retry-unfilled"])
    cli.main()
    assert len(calls) == 1


@pytest.mark.parametrize("text,deadline,why", [
    ("2026-2027 tuition is $1,200 per credit.", False, None),
    ("Tuition for 2026-27: $1,200.", False, None),
    ("Tuition is $1,200 per credit.", False, None),
    ("2025-2026 tuition is $1,150.", False, "names academic year 2025-2026"),
    ("2027-28 rates: $1,300.", False, "names academic year 2027-2028"),
    ("Fall 2026 deadline: June 1, 2026.", True, None),
    ("Spring 2027 start: apply by November 15, 2026.", True, None),
    ("Fall 2027 priority deadline: January 15, 2027.", True, "names term Fall 2027"),
    ("Summer 2026 start.", False, "names term Summer 2026"),
    ("Round 1: October 15, 2025.", True, None),          # deadlines precede the cycle
    ("Average GMAT 640 (Fall 2024 profile).", False, "names term Fall 2024"),
    ("Class of 2025 average GPA 3.4.", False, "names class of 2025"),
    ("Class of 2028 average GPA 3.4.", False, None),
    ("Ranked #5 in 2025 U.S. News.", False, "names year 2025"),
    ("Ranked #5 in 2026 Best Online Programs.", False, None),
])
def test_only_the_2026_2027_cycle_passes(text, deadline, why):
    ctx = agent.Context(date(2026, 9, 24), cli.DEFAULT_CYCLE, "Online MBA")
    assert ctx.academic_years == (2026, 2027)
    assert ctx.cycle_violation(text, deadline) == why


def test_off_cycle_quote_is_rejected_even_if_the_model_says_exact(monkeypatch):
    ctx = agent.Context(date(2026, 9, 24), cli.DEFAULT_CYCLE, "Online MBA")
    page = "Tuition 2025-2026: $1,150 per credit. Credits: 36 total."
    script(monkeypatch, {"https://x.edu/": (page, [])}, {"https://x.edu/": {"found": [
        finding("Tuition", "$1,150", "Tuition 2025-2026: $1,150 per credit."),
        finding("Credits", "36", "Credits: 36 total.", cycle="unstated")], "next_urls": []}})
    results, meta = agent.research_college("X", ["https://x.edu/"], ["Tuition", "Credits"], ctx)
    assert results["Tuition"]["status"] == "not-found"
    assert meta["rejected"] == {"off-cycle (names academic year 2025-2026)": 1}


def test_claude_takes_over_what_rcac_missed_and_each_model_verifies_its_own_values(monkeypatch):
    ctx = agent.Context(date(2026, 9, 24), cli.DEFAULT_CYCLE, "Online MBA")
    page = "Fall 2026 tuition: $1,200 per credit. The program is 36 credit hours. Deposit $500."
    monkeypatch.setattr(agent, "fetch", lambda url: Page(page + PAD, [], ""))
    asked = []

    def model(name, found):
        def call(prompt):
            asked.append((name, "VERIFY" if "Check one value" in prompt else "EXTRACT"))
            if "Check one value" in prompt:
                bad = "Extracted value: $999" in prompt
                return json.dumps({"verdict": "incorrect" if bad else "correct", "failed_check": 2 if bad else None,
                                   "corrected_value": "", "reason": "r"})
            return json.dumps({"found": found, "next_urls": []})
        return call

    monkeypatch.setattr(agent, "call_rcac", model("rcac", [
        finding("Tuition", "$1,200", "Fall 2026 tuition: $1,200 per credit.")]))
    monkeypatch.setattr(agent, "call_claude", model("claude", [
        finding("Credits", "36", "The program is 36 credit hours.", cycle="exact"),
        finding("Deposit", "$999", "Deposit $500.")]))  # digits not in quote -> rejected in code
    results, meta = agent.research_college("X", ["https://x.edu/"], ["Tuition", "Credits", "Deposit"], ctx,
                                           models=("rcac", "claude"), verify=True)
    assert (results["Tuition"]["value"], results["Tuition"]["verified_by"]) == ("$1,200", "rcac")
    assert (results["Credits"]["value"], results["Credits"]["verified_by"]) == ("36", "claude")
    assert results["Deposit"]["status"] == "not-found"
    assert ("rcac", "VERIFY") in asked and ("claude", "VERIFY") in asked
    assert asked.index(("claude", "EXTRACT")) > asked.index(("rcac", "VERIFY"))  # Claude only after RCAC finished
    assert meta["models"] == {"rcac": "ok", "claude": "ok"}


def test_verifier_rejection_drops_the_value_and_rcac_outage_hands_over_to_claude(monkeypatch):
    ctx = agent.Context(date(2026, 9, 24), cli.DEFAULT_CYCLE, "Online MBA")
    page = "Application fee $75 for international applicants. Application fee $60."
    monkeypatch.setattr(agent, "fetch", lambda url: Page(page + PAD, [], ""))

    def down(prompt):
        raise RuntimeError("RCAC API failed after retries")

    def claude_call(prompt):
        if "Check one value" in prompt:
            bad = "Extracted value: $75" in prompt
            return json.dumps({"verdict": "incorrect" if bad else "correct", "reason": "international only" if bad else "ok"})
        if "Pick the candidate" in prompt:
            return json.dumps({"choice": 1, "reason": "r"})
        return json.dumps({"found": [finding("Fee", "$75", "Application fee $75 for international applicants."),
                                     finding("Fee", "$60", "Application fee $60.")], "next_urls": []})

    monkeypatch.setattr(agent, "call_rcac", down)
    monkeypatch.setattr(agent, "call_claude", claude_call)
    results, meta = agent.research_college("X", ["https://x.edu/"], ["Fee"], ctx, models=("rcac", "claude"), verify=True)
    assert meta["models"]["rcac"].startswith("down")
    assert (results["Fee"]["value"], results["Fee"]["status"], results["Fee"]["verified_by"]) == ("$60", "single-source", "claude")
    assert [c["value"] for c in results["Fee"]["verify_rejected"]] == ["$75"]
    assert agent.best_effort(results["Fee"])[0] == "$60"


def test_field_guide_is_loaded_into_prompts_and_link_keywords(monkeypatch):
    assert "Fall 2026" in skills.cycle_rules()
    assert "per credit hour for in-state" in skills.field_hint("In-State Tuition Per Credit")
    assert "out-of-state" in skills.field_hint("Out-of-State Tuition Per Credit").casefold()
    assert "Yes" in skills.field_hint("Deposit Required?") and "$500" in skills.field_hint("Deposit")
    assert "three-year" in skills.field_hint("Accept 3yr-degree from India? (Y/N)")
    assert {"tuition", "bursar"} <= skills.field_keywords(["International Tuition Total", "In-State Tuition Per Credit"])
    ctx = agent.Context(date(2026, 9, 24), cli.DEFAULT_CYCLE, "Online MBA")
    prompts = script(monkeypatch, {"https://x.edu/": ("x", [])}, {"https://x.edu/": {"found": [], "next_urls": []}})
    agent.research_college("X", ["https://x.edu/"], ["Average GMAT"], ctx)
    assert "Average (mean) GMAT" in prompts[0] and "2026-2027 academic year ONLY" in prompts[0]
